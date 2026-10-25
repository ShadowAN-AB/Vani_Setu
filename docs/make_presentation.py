#!/usr/bin/env python3
"""Build a 12-slide 16:9 PowerPoint for the Vani-Setu viva / demo."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# Brand (matches the web app)
TEAL = RGBColor(0x0B, 0x5F, 0x66)
TEAL_DEEP = RGBColor(0x08, 0x48, 0x4E)
TEAL_SOFT = RGBColor(0xD7, 0xEC, 0xEE)
PAPER = RGBColor(0xF3, 0xF7, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x10, 0x29, 0x2D)
INK_SOFT = RGBColor(0x3D, 0x4F, 0x52)
MUTED = RGBColor(0x6B, 0x7C, 0x7F)
AMBER = RGBColor(0xC9, 0x78, 0x16)
AMBER_SOFT = RGBColor(0xF8, 0xE8, 0xD2)
LINE = RGBColor(0xD5, 0xE0, 0xE2)
SUCCESS = RGBColor(0x1F, 0x7A, 0x4D)

W, H = Inches(13.333), Inches(7.5)
OUT = Path(__file__).resolve().parent / "Vani_Setu_Presentation.pptx"


def set_run(run, text, size=18, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
    return sh


def add_round(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
    # slightly tighter corners
    try:
        adj = sh.adjustments
        adj[0] = 0.12
    except Exception:
        pass
    return sh


def add_tb(slide, l, t, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size=size, bold=bold, color=color, font=font)
    return box


def add_para(tf, text, size=16, bold=False, color=INK, space_before=6, space_after=0, align=PP_ALIGN.LEFT):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    set_run(run, text, size=size, bold=bold, color=color)
    return p


def footer(slide, n, total=12):
    add_rect(slide, Inches(0), Inches(7.28), W, Inches(0.22), TEAL)
    add_tb(slide, Inches(0.5), Inches(7.28), Inches(10), Inches(0.22),
           "Vani-Setu  ·  voice-bridge  ·  vani-setu.onrender.com",
           size=10, color=WHITE)
    add_tb(slide, Inches(11.4), Inches(7.28), Inches(1.5), Inches(0.22),
           f"{n}  /  {total}", size=10, color=WHITE, align=PP_ALIGN.RIGHT)


def kicker(slide, text):
    add_tb(slide, Inches(0.55), Inches(0.28), Inches(8), Inches(0.32),
           text.upper(), size=12, bold=True, color=TEAL, font="Calibri")


def title(slide, text):
    add_tb(slide, Inches(0.5), Inches(0.5), Inches(12.3), Inches(0.6),
           text, size=32, bold=True, color=TEAL_DEEP, font="Georgia")


def card(slide, l, t, w, h, head, body, accent=TEAL):
    add_round(slide, l, t, w, h, WHITE)
    add_rect(slide, l, t, Inches(0.08), h, accent)
    box = slide.shapes.add_textbox(l + Inches(0.22), t + Inches(0.14), w - Inches(0.34), h - Inches(0.22))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, head, size=16, bold=True, color=TEAL_DEEP)
    add_para(tf, body, size=13, color=INK_SOFT, space_before=6)


def make():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # ---------- 1 Cover ----------
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, TEAL_DEEP)
    add_rect(s, 0, 0, Inches(0.22), H, AMBER)
    add_tb(s, Inches(0.8), Inches(1.55), Inches(11), Inches(0.4),
           "COLLEGE PROJECT  ·  SIGN LANGUAGE TRANSLATOR", size=14, bold=True, color=AMBER_SOFT)
    add_tb(s, Inches(0.75), Inches(2.05), Inches(12), Inches(1.15),
           "Vani-Setu", size=66, bold=True, color=WHITE, font="Georgia")
    add_tb(s, Inches(0.8), Inches(3.25), Inches(11.5), Inches(0.7),
           "voice-bridge  —  real-time ASL and ISL from a laptop camera",
           size=22, color=TEAL_SOFT)
    add_round(s, Inches(0.8), Inches(4.25), Inches(5.6), Inches(0.55), TEAL)
    add_tb(s, Inches(0.9), Inches(4.32), Inches(5.4), Inches(0.45),
           "  Live   vani-setu.onrender.com", size=16, bold=True, color=WHITE)
    add_tb(s, Inches(0.8), Inches(5.15), Inches(11), Inches(0.8),
           "Camera  →  MediaPipe landmarks  →  SVM / pose rules  →  sentence, speech, 25+ languages",
           size=16, color=TEAL_SOFT)
    add_tb(s, Inches(0.8), Inches(6.65), Inches(11), Inches(0.35),
           "12 slides  ·  product  ·  stack  ·  features  ·  demo", size=13, color=MUTED)

    # ---------- 2 Problem ----------
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, PAPER)
    footer(s, 2)
    kicker(s, "Why this exists")
    title(s, "The communication gap")
    points = [
        ("Independence", "Deaf and hard-of-hearing people often need a human interpreter for everyday talk. A laptop camera can sit in between."),
        ("Classroom need", "A college demo that is honest: one hand, live video, letters and short phrases — not a fake ‘full language’ claim."),
        ("Local first", "Works on localhost for a viva. Same app is also live on HTTPS so a judge can open a link."),
    ]
    for i, (h, b) in enumerate(points):
        card(s, Inches(0.5), Inches(1.35) + Inches(i * 1.75), Inches(12.3), Inches(1.6), h, b)
    add_tb(s, Inches(0.7), Inches(6.65), Inches(12), Inches(0.4),
           "Name: Vani-Setu  —  Sanskrit-inspired “voice-bridge”.", size=14, color=MUTED)

    # ---------- 3 Product ----------
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, PAPER)
    footer(s, 3)
    kicker(s, "The product")
    title(s, "What Vani-Setu is")
    add_tb(s, Inches(0.55), Inches(1.2), Inches(12.2), Inches(0.55),
           "A Flask web app. The camera stays in the browser. The server classifies, translates, and stores accounts.",
           size=16, color=INK_SOFT)
    tiles = [
        ("01", "See", "Laptop webcam + MediaPipe tracks 21 points on one hand. Tracking never leaves the device."),
        ("02", "Read", "Spell mode: A–Z via an RBF SVM. Phrase / ISL: classroom poses via finger-geometry rules."),
        ("03", "Speak", "Build a sentence, hear it (Web Speech API), translate into 25+ languages. Hindi has an offline fallback."),
        ("04", "Learn", "Practice mode scores letters. Record samples and retrain the SVM from the browser."),
    ]
    for i, (n, h, b) in enumerate(tiles):
        col, row = i % 2, i // 2
        l = Inches(0.5) + col * Inches(6.4)
        t = Inches(1.85) + row * Inches(2.4)
        add_round(s, l, t, Inches(6.15), Inches(2.2), WHITE)
        add_tb(s, l + Inches(0.3), t + Inches(0.25), Inches(1.2), Inches(0.4), n, size=20, bold=True, color=AMBER)
        add_tb(s, l + Inches(0.3), t + Inches(0.7), Inches(5.5), Inches(0.4), h, size=22, bold=True, color=TEAL_DEEP, font="Georgia")
        add_tb(s, l + Inches(0.3), t + Inches(1.15), Inches(5.5), Inches(0.85), b, size=14, color=INK_SOFT)

    # ---------- 4 Features ----------
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, PAPER)
    footer(s, 4)
    kicker(s, "In brief")
    title(s, "Features")
    feats = [
        ("Spell A–Z", "SVM on 63-D landmarks, ~95.1% held-out"),
        ("Phrase mode", "Hello, Yes, OK, Peace, I Love You…"),
        ("ISL classroom set", "Namaste, Thank You, Water, Help…"),
        ("Confidence bar", "Live probability while spelling"),
        ("Word chips", "Tap a suggestion to finish a word"),
        ("Practice mode", "Prompted letters + score / streak"),
        ("Translate + Speak", "25+ languages, Hindi offline backup"),
        ("Sign Guide", "On-screen pose hints for the demo"),
        ("Record + Retrain", "Collect samples; hot-reload the model"),
        ("Password login", "Hashed accounts in SQLite"),
        ("Guest login", "Skip accounts on the public URL"),
        ("Health check", "/api/health before a viva"),
    ]
    for i, (h, b) in enumerate(feats):
        col, row = i % 4, i // 4
        l = Inches(0.45) + col * Inches(3.2)
        t = Inches(1.3) + row * Inches(1.85)
        add_round(s, l, t, Inches(3.05), Inches(1.7), WHITE)
        add_rect(s, l, t, Inches(3.05), Inches(0.08), TEAL if row == 0 else AMBER if row == 1 else SUCCESS)
        add_tb(s, l + Inches(0.18), t + Inches(0.28), Inches(2.7), Inches(0.5), h, size=15, bold=True, color=TEAL_DEEP)
        add_tb(s, l + Inches(0.18), t + Inches(0.8), Inches(2.7), Inches(0.7), b, size=12, color=INK_SOFT)

    # ---------- 5 Modes ----------
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, PAPER)
    footer(s, 5)
    kicker(s, "Input")
    title(s, "Three modes")
    modes = [
        (TEAL, "SPELL", "Machine learning",
         "Fingerspelling A–Z.\nRBF SVM on 2,122 samples.\n0.50 confidence floor +\nmulti-frame voting."),
        (AMBER, "PHRASE", "Rule-based",
         "Whole words from one-hand geometry.\nHello, Stop, Good, Yes, No,\nI Love You, Peace, OK…"),
        (SUCCESS, "ISL", "Classroom set",
         "Simplified single-hand signs\nfor teaching — not a certified\nISL lexicon. Switch US / IN."),
    ]
    for i, (c, name, kind, body) in enumerate(modes):
        l = Inches(0.5) + i * Inches(4.2)
        add_round(s, l, Inches(1.35), Inches(3.95), Inches(5.35), WHITE)
        add_rect(s, l, Inches(1.35), Inches(3.95), Inches(0.14), c)
        add_tb(s, l + Inches(0.3), Inches(1.7), Inches(3.35), Inches(0.45), name, size=22, bold=True, color=c)
        add_tb(s, l + Inches(0.3), Inches(2.2), Inches(3.35), Inches(0.4), kind, size=14, bold=True, color=MUTED)
        add_tb(s, l + Inches(0.3), Inches(2.8), Inches(3.35), Inches(3.4), body, size=16, color=INK)

    # ---------- 6 Tech stack ----------
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, PAPER)
    footer(s, 6)
    kicker(s, "Tools")
    title(s, "Tech stack")
    stacks = [
        ("Interface", "HTML5  ·  CSS3  ·  vanilla JS\nGoogle Fonts (Fraunces / Sora)"),
        ("Vision", "MediaPipe Hand Landmarker\ngetUserMedia camera in-browser"),
        ("Backend", "Python 3  ·  Flask  ·  Flask-CORS\nGunicorn on the host"),
        ("Machine learning", "scikit-learn RBF SVM  ·  joblib\npandas  ·  matplotlib reports"),
        ("Accounts", "Flask-SQLAlchemy  ·  Werkzeug\nSQLite  ·  optional Postgres"),
        ("Speech & language", "Web Speech API\ndeep-translator  ·  Hindi fallback"),
        ("Data", "ASL Now landmarks (MIT)\n63-D wrist-relative features"),
        ("Ship", "run.sh on localhost\nRender HTTPS  ·  GitHub"),
    ]
    for i, (h, b) in enumerate(stacks):
        col, row = i % 4, i // 4
        l = Inches(0.4) + col * Inches(3.2)
        t = Inches(1.3) + row * Inches(2.7)
        add_round(s, l, t, Inches(3.05), Inches(2.45), WHITE)
        add_rect(s, l, t, Inches(3.05), Inches(0.5), TEAL if row == 0 else TEAL_DEEP)
        add_tb(s, l + Inches(0.15), t + Inches(0.08), Inches(2.75), Inches(0.4), h, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_tb(s, l + Inches(0.18), t + Inches(0.7), Inches(2.7), Inches(1.55), b, size=13, color=INK, align=PP_ALIGN.CENTER)

    # ---------- 7 Architecture ----------
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, PAPER)
    footer(s, 7)
    kicker(s, "Design")
    title(s, "How the pieces connect")
    steps = [
        ("Browser", "HTML / CSS / JS\nMediaPipe 21 pts"),
        ("Flask API", "/api/classify\n/api/translate\n/api/login"),
        ("Models", "SVM  ·  Spell\nRules  ·  Phrase"),
        ("Store", "SQLite users\njoblib model"),
    ]
    for i, (h, b) in enumerate(steps):
        l = Inches(0.45) + i * Inches(3.2)
        add_round(s, l, Inches(1.45), Inches(2.85), Inches(2.35), WHITE)
        add_tb(s, l + Inches(0.15), Inches(1.6), Inches(2.55), Inches(0.45), h, size=18, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
        add_tb(s, l + Inches(0.15), Inches(2.15), Inches(2.55), Inches(1.4), b, size=14, color=INK_SOFT, align=PP_ALIGN.CENTER)
        if i < 3:
            add_tb(s, l + Inches(2.7), Inches(2.2), Inches(0.5), Inches(0.5), "→", size=28, bold=True, color=AMBER)
    notes = [
        "Landmarks are computed on-device. Only a 63-D vector is posted for Spell.",
        "Phrase / ISL stay rule-based unless you record phrase samples.",
        "Session after login is sessionStorage — a demo gate, not bank-grade auth.",
        "Public camera needs HTTPS. That is why the live site is on Render.",
    ]
    for i, n in enumerate(notes):
        add_round(s, Inches(0.45), Inches(4.1) + Inches(i * 0.7), Inches(12.4), Inches(0.62), WHITE)
        add_tb(s, Inches(0.7), Inches(4.18) + Inches(i * 0.7), Inches(12), Inches(0.5), "▸  " + n, size=15, color=INK)

    # ---------- 8 ML ----------
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, PAPER)
    footer(s, 8)
    kicker(s, "Why this classifier")
    title(s, "Spell-mode ML")
    add_tb(s, Inches(0.55), Inches(1.2), Inches(12), Inches(0.5),
           "Same 63-D features, five models. SVM won on a held-out split — not by preference.",
           size=16, color=INK_SOFT)
    rows = [
        ("Model", "Test accuracy", "5-fold CV", True),
        ("SVM (RBF)", "95.1%", "94.1%", True),
        ("Logistic Regression", "94.0%", "92.5%", False),
        ("Random Forest", "88.1%", "87.9%", False),
        ("K-Nearest Neighbours", "85.7%", "84.9%", False),
        ("Decision Tree", "74.6%", "74.8%", False),
    ]
    for i, (a, b, c, hi) in enumerate(rows):
        t = Inches(1.8) + Inches(i * 0.58)
        bg = TEAL if i == 0 else (TEAL_SOFT if hi and i == 1 else WHITE)
        fg = WHITE if i == 0 else (TEAL_DEEP if hi and i == 1 else INK)
        add_rect(s, Inches(0.55), t, Inches(8.3), Inches(0.55), bg)
        add_tb(s, Inches(0.7), t + Inches(0.08), Inches(3.8), Inches(0.42), a, size=15, bold=(i == 0 or hi), color=fg)
        add_tb(s, Inches(4.6), t + Inches(0.08), Inches(2.0), Inches(0.42), b, size=15, bold=(i == 0 or hi), color=fg)
        add_tb(s, Inches(6.8), t + Inches(0.08), Inches(1.9), Inches(0.42), c, size=15, bold=(i == 0 or hi), color=fg)
    add_round(s, Inches(9.1), Inches(1.8), Inches(3.7), Inches(4.7), WHITE)
    add_tb(s, Inches(9.3), Inches(2.0), Inches(3.3), Inches(0.4), "Live filters", size=16, bold=True, color=TEAL)
    extras = [
        "2,122 letter samples, A–Z",
        "Confidence floor 0.50",
        "Keeps ~96% of correct calls",
        "Rejects ~69% of errors",
        "Multi-frame voting on top",
        "Dataset: ASL Now (MIT)",
    ]
    for i, e in enumerate(extras):
        add_tb(s, Inches(9.3), Inches(2.5) + Inches(i * 0.55), Inches(3.3), Inches(0.5), "•  " + e, size=14, color=INK)

    # ---------- 9 Auth ----------
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, PAPER)
    footer(s, 9)
    kicker(s, "Accounts & shipping")
    title(s, "Login and deploy")
    left = [
        ("Localhost (viva)", "AUTH_MODE=password\nEmail + hashed password\nSQLite file users.db\nhttp://127.0.0.1:5001\n./run.sh  —  port 5001"),
        ("Public demo", "AUTH_MODE=guest\nContinue as guest\nHTTPS on Render\nvani-setu.onrender.com\nGunicorn  ·  0.0.0.0"),
    ]
    for i, (h, b) in enumerate(left):
        l = Inches(0.5) + i * Inches(6.4)
        add_round(s, l, Inches(1.35), Inches(6.1), Inches(3.35), WHITE)
        add_rect(s, l, Inches(1.35), Inches(6.1), Inches(0.55), TEAL if i == 0 else AMBER)
        add_tb(s, l + Inches(0.25), Inches(1.42), Inches(5.6), Inches(0.42), h, size=18, bold=True, color=WHITE)
        add_tb(s, l + Inches(0.3), Inches(2.1), Inches(5.5), Inches(2.4), b, size=16, color=INK)
    add_round(s, Inches(0.5), Inches(4.9), Inches(12.3), Inches(2.0), WHITE)
    add_tb(s, Inches(0.75), Inches(5.1), Inches(11.8), Inches(0.4), "No MongoDB. No paid SMS. Optional Postgres via DATABASE_URL.", size=16, bold=True, color=TEAL_DEEP)
    add_tb(s, Inches(0.75), Inches(5.55), Inches(11.8), Inches(1.1),
           "Free Render sleeps after idle time — first open can take ~1 minute. Letter CSV is not on the host, so Retrain belongs on localhost. Spell still uses the shipped SVM.",
           size=15, color=INK_SOFT)

    # ---------- 10 Demo ----------
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, PAPER)
    footer(s, 10)
    kicker(s, "Show this")
    title(s, "Demo path")
    steps = [
        ("1", "Open", "vani-setu.onrender.com\nor localhost:5001"),
        ("2", "Enter", "Continue as guest\nor Create account"),
        ("3", "Camera", "Start camera\nallow the browser"),
        ("4", "Spell", "Hold L, V, W, B\nwatch confidence"),
        ("5", "Phrase", "Hello / Peace / OK\nthen Speak"),
        ("6", "Extra", "Word chips, Practice,\nSign Guide, Logout"),
    ]
    for i, (n, h, b) in enumerate(steps):
        col, row = i % 3, i // 3
        l = Inches(0.5) + col * Inches(4.2)
        t = Inches(1.35) + row * Inches(2.7)
        add_round(s, l, t, Inches(3.95), Inches(2.45), WHITE)
        add_round(s, l + Inches(0.25), t + Inches(0.25), Inches(0.55), Inches(0.55), TEAL)
        add_tb(s, l + Inches(0.25), t + Inches(0.3), Inches(0.55), Inches(0.45), n, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_tb(s, l + Inches(0.95), t + Inches(0.3), Inches(2.7), Inches(0.45), h, size=20, bold=True, color=TEAL_DEEP)
        add_tb(s, l + Inches(0.3), t + Inches(1.05), Inches(3.35), Inches(1.15), b, size=15, color=INK_SOFT)

    # ---------- 11 Limits ----------
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, PAPER)
    footer(s, 11)
    kicker(s, "Say this in viva")
    title(s, "Honest limits")
    limits = [
        ("One hand, mostly static poses", "Motion letters J and Z are weak. No two-handed ISL."),
        ("Fingerspelling ≠ full language", "Spell ML is ASL letters, not a sign-language corpus."),
        ("ISL is a teaching set", "Classroom approximation, not a certified lexicon."),
        ("Phrase is rule-based", "Until you record landmarks_phrases.csv."),
        ("Camera needs HTTPS off localhost", "Browsers block webcam on public http://."),
        ("Demo accounts, not production IAM", "sessionStorage after login; Render disk is ephemeral."),
    ]
    for i, (h, b) in enumerate(limits):
        col, row = i % 2, i // 2
        l = Inches(0.45) + col * Inches(6.4)
        t = Inches(1.3) + row * Inches(1.8)
        add_round(s, l, t, Inches(6.15), Inches(1.65), WHITE)
        add_tb(s, l + Inches(0.3), t + Inches(0.25), Inches(5.55), Inches(0.45), h, size=16, bold=True, color=TEAL_DEEP)
        add_tb(s, l + Inches(0.3), t + Inches(0.75), Inches(5.55), Inches(0.7), b, size=14, color=INK_SOFT)

    # ---------- 12 Close ----------
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, TEAL_DEEP)
    add_rect(s, 0, 0, Inches(0.22), H, AMBER)
    add_tb(s, Inches(0.8), Inches(1.7), Inches(12), Inches(0.9),
           "Thank you", size=52, bold=True, color=WHITE, font="Georgia")
    add_tb(s, Inches(0.8), Inches(2.7), Inches(12), Inches(0.5),
           "Questions?", size=24, color=TEAL_SOFT)
    links = [
        "Live     vani-setu.onrender.com",
        "Code     github.com/ShadowAN-AB/Vani_Setu",
        "Local    ./run.sh   →   http://127.0.0.1:5001",
    ]
    for i, line in enumerate(links):
        add_tb(s, Inches(0.8), Inches(3.55) + Inches(i * 0.5), Inches(11), Inches(0.45), line, size=18, color=WHITE)
    add_tb(s, Inches(0.8), Inches(5.5), Inches(11.5), Inches(1.1),
           "Vani-Setu  ·  MediaPipe + Flask + RBF SVM  ·  95.1% letters  ·  one honest classroom translator",
           size=15, color=TEAL_SOFT)
    add_tb(s, Inches(0.8), Inches(6.7), Inches(11), Inches(0.35), "12 / 12", size=12, color=MUTED)

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    make()
